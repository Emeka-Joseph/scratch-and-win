"""CSV / Excel export helpers used by the superadmin and brand dashboards."""
import csv
import io
from datetime import datetime

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_response(fmt, filename, headers, rows, sheet_title="Export"):
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    if fmt == "xlsx":
        return _xlsx(f"{filename}-{stamp}.xlsx", headers, rows, sheet_title)
    return _csv(f"{filename}-{stamp}.csv", headers, rows)


def _csv(name, headers, rows):
    buf = io.StringIO()
    buf.write("﻿")  # BOM so Excel opens Arabic text correctly
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return Response(buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _xlsx(name, headers, rows, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="0B0B0B")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    widths = [len(str(h)) for h in headers]
    for row in rows:
        row = ["" if v is None else v for v in row]
        ws.append(row)
        for i, v in enumerate(row):
            widths[i] = min(max(widths[i], len(str(v))), 50)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 3
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
